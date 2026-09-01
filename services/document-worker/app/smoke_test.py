from __future__ import annotations

import hashlib
import io
import json
import urllib.request
import uuid

import fitz
from PIL import Image, ImageDraw, ImageFont
from openpyxl import Workbook

from app.main import extract_csv, extract_image, extract_pdf, extract_xlsx, parse_docling_result


def test_ocr_image() -> None:
    image = Image.new("RGB", (1500, 320), "white")
    draw = ImageDraw.Draw(image)
    font = ImageFont.truetype("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf", 72)
    draw.text((40, 80), "META 100 REALIZADO 75", fill="black", font=font)
    buffer = io.BytesIO()
    image.save(buffer, format="JPEG", quality=95)
    content = buffer.getvalue()
    digest = f"sha256:{hashlib.sha256(content).hexdigest()}"
    result = extract_image(content, "ocr-smoke.jpg", "image/jpeg", digest)
    normalized = result.text.upper()
    assert "META" in normalized and "100" in normalized and "75" in normalized, result.text
    assert result.evidence[0].locator == "image:1"


def test_native_pdf() -> None:
    document = fitz.open()
    page = document.new_page()
    for index in range(5):
        page.insert_text((72, 80 + index * 28), f"INDICADOR CONSORCIO META 100 REALIZADO 80 PERIODO AGOSTO LINHA {index + 1}", fontsize=11)
    content = document.tobytes()
    digest = f"sha256:{hashlib.sha256(content).hexdigest()}"
    result = extract_pdf(content, "native-smoke.pdf", digest)
    assert result.extraction_method == "PDF_NATIVE_TEXT", result.extraction_method
    assert "CONSORCIO" in result.text
    assert result.evidence[0].locator == "page:1"


def test_docling_structured_table() -> None:
    payload = {"version": "1.30.0", "document": {"md_content": "# POBJ", "json_content": {
        "tables": [{"prov": [{"page_no": 2}], "data": {"num_rows": 2, "num_cols": 3, "table_cells": [
            {"start_row_offset_idx": 0, "end_row_offset_idx": 1, "start_col_offset_idx": 0, "end_col_offset_idx": 1, "text": "INDICADOR"},
            {"start_row_offset_idx": 0, "end_row_offset_idx": 1, "start_col_offset_idx": 1, "end_col_offset_idx": 2, "text": "META"},
            {"start_row_offset_idx": 0, "end_row_offset_idx": 1, "start_col_offset_idx": 2, "end_col_offset_idx": 3, "text": "REALIZADO"},
            {"start_row_offset_idx": 1, "end_row_offset_idx": 2, "start_col_offset_idx": 0, "end_col_offset_idx": 1, "text": "Consórcio"},
            {"start_row_offset_idx": 1, "end_row_offset_idx": 2, "start_col_offset_idx": 1, "end_col_offset_idx": 2, "text": "100"},
            {"start_row_offset_idx": 1, "end_row_offset_idx": 2, "start_col_offset_idx": 2, "end_col_offset_idx": 3, "text": "80"},
        ]}}], "texts": [{"label": "title", "text": "POBJ Agosto", "prov": [{"page_no": 1}]}]
    }}}
    result = parse_docling_result(payload, "pobj.pdf", "application/pdf", "sha256:" + "0" * 64, "success", 1200)
    assert result.extraction_method == "DOCLING_TABLEFORMER"
    assert result.tables[0].headers == ["INDICADOR", "META", "REALIZADO"]
    assert result.tables[0].rows[0] == ["Consórcio", "100", "80"]
    assert result.tables[0].page_number == 2


def test_native_tabular_formats() -> None:
    csv_content = b"INDICADOR,META,REALIZADO\nConsorcio,100,80\n"
    csv_result = extract_csv(csv_content, "pobj.csv", f"sha256:{hashlib.sha256(csv_content).hexdigest()}")
    assert csv_result.extraction_method == "CSV_NATIVE"
    assert csv_result.tables[0].headers == ["INDICADOR", "META", "REALIZADO"]
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["INDICADOR", "META", "REALIZADO"])
    sheet.append(["Consórcio", 100, 80])
    buffer = io.BytesIO()
    workbook.save(buffer)
    xlsx_content = buffer.getvalue()
    xlsx_result = extract_xlsx(xlsx_content, "pobj.xlsx", f"sha256:{hashlib.sha256(xlsx_content).hexdigest()}")
    assert xlsx_result.extraction_method == "XLSX_NATIVE"
    assert xlsx_result.tables[0].rows[0] == ["Consórcio", "100", "80"]


def test_http_endpoint() -> None:
    image = Image.new("RGB", (1200, 260), "white")
    draw = ImageDraw.Draw(image)
    font = ImageFont.truetype("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf", 58)
    draw.text((30, 70), "PONTUACAO 42 META 100", fill="black", font=font)
    image_buffer = io.BytesIO()
    image.save(image_buffer, format="JPEG", quality=95)
    content = image_buffer.getvalue()
    digest = f"sha256:{hashlib.sha256(content).hexdigest()}"
    job_id = "smoke-http-job"
    metadata = json.dumps({
        "job_id": job_id,
        "document": {"content_hash": digest, "file_name": "http-smoke.jpg"},
        "security": {"external_effects_allowed": False},
    })
    boundary = f"----diretor360{uuid.uuid4().hex}"
    body = (
        f"--{boundary}\r\nContent-Disposition: form-data; name=\"metadata\"\r\n\r\n{metadata}\r\n"
        f"--{boundary}\r\nContent-Disposition: form-data; name=\"document\"; filename=\"http-smoke.jpg\"\r\n"
        "Content-Type: image/jpeg\r\n\r\n"
    ).encode() + content + f"\r\n--{boundary}--\r\n".encode()
    request = urllib.request.Request(
        "http://127.0.0.1:8787/v1/process",
        data=body,
        headers={
            "Content-Type": f"multipart/form-data; boundary={boundary}",
            "X-Job-Id": job_id,
            "X-Content-Trust": "UNTRUSTED",
        },
        method="POST",
    )
    with urllib.request.urlopen(request, timeout=360) as response:
        result = json.load(response)
    assert result["ok"] is True
    assert result["job_id"] == job_id
    assert "PONTUACAO" in result["extraction"]["text"].upper()


if __name__ == "__main__":
    test_ocr_image()
    test_native_pdf()
    test_docling_structured_table()
    test_native_tabular_formats()
    test_http_endpoint()
    print("DOCUMENT_WORKER_EXTRACTION: PASS")
